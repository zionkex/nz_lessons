from seleniumbase import Driver
import customtkinter as ctk
from customtkinter import filedialog
from docx import Document


class WebAutomationApp():
    def __init__(self):
            self.root = ctk.CTk()
            self.driver = Driver(undetectable=True)
            self.root.title('NZ.UA')
            self.root.geometry('400x200+300+100')
            
            self.root_label = ctk.CTkLabel(self.root, text='Логін:', font=('Arial', 18, 'bold'))
            self.root_label.grid(row=4, column=0, padx=10, pady=10)

            self.root_entry = ctk.CTkEntry(self.root, width= 200, font=('Arial', 18))
            self.root_entry.grid(row=4, column=1, padx=10, pady=10, sticky='w')

            self.pass_label = ctk.CTkLabel(self.root, text='Пароль:', font=('Arial', 18, 'bold'))
            self.pass_label.grid(row=5, column=0, padx=10, pady=10)

            self.pass_entry = ctk.CTkEntry(self.root, show='*', width= 200, font=('Arial', 18))
            self.pass_entry.grid(row=5, column=1, padx=10, pady=10, sticky='w')

            self.root_start = ctk.CTkButton(self.root, text='Увійти', font=('Arial', 18, 'bold'), command=self.login)
            self.root_start.grid(row=6, column=1)
            self.root.grid_columnconfigure(1, minsize=100)
            
            self.root.mainloop()

    def login(self):
            self.root.withdraw()
            self.driver.get('https://nz.ua/')
            self.driver.click('button:contains("Увійти")')
            self.driver.type('#loginform-login', self.root_entry.get())
            self.driver.sleep(1)
            self.driver.type('#loginform-password', self.pass_entry.get())
            self.driver.sleep(1)
            self.driver.click('button.form-submit-btn')
            self.driver.click('a:contains("Журнали")')
            self.open_second_window()
            # self.driver.wait_for_element('div.dropdown pull-right',timeout=500)
    def connect_journal(self, combo, combo2):
        desired_subject = combo.get()
        desired_class = combo2.get()
        subject_elements = self.driver.find_elements('xpath', "//td[not(a[contains(@class, 'gray-button-2')])]")
        for element in subject_elements:
            subject_name = element.text.strip() 
            if subject_name == desired_subject:
                class_links = element.find_elements('xpath', "./following-sibling::td/a[@class='gray-button-2']")
                for class_link in class_links:
                    class_text = class_link.text
                    if desired_class in class_text:
                        class_link.click()
                        self.hide_widgets()
                        break
                break
    def update_class_list(self, selected_value):
        selected_subject = selected_value
        class_row = self.driver.find_elements('xpath', f"//td[normalize-space()='{selected_subject}']/following-sibling::td/a[@class='gray-button-2']")
        class_list = [element.text.strip() for element in class_row if element.text.strip()]
        self.combo2.configure(values=class_list)
        self.combo2.set('Виберіть клас')
    def journal(self):
        elements = self.driver.find_elements('xpath', "//table[@class='journal-choose']//td[not(a[contains(@class, 'gray-button-2')])]")
        subjects = [element.text.strip() for element in elements if element.text.strip()]
        return subjects
    

    def open_second_window(self):
        self.root.withdraw()
        self.second_window = ctk.CTkToplevel(self.root)
        self.second_window.geometry('380x150+200+100')
        self.second_window.title('NZ.UA')
        self.back_button = ctk.CTkButton(self.second_window, width=80,text='Вихід', command=self.back_journals , font=('Arial', 16, 'bold'))
        self.back_button.place(x=17, y=250)
        # Initialize combo boxes and configure their properties
        self.combo_values = self.journal()
        self.combo = ctk.CTkComboBox(self.second_window, width=150, values=self.combo_values, command=lambda value: self.update_class_list(value), font=('Arial', 18, 'bold'))
        self.combo.set('Виберіть предмет')
        self.combo.place(x=17, y=25)

        self.combo2 = ctk.CTkComboBox(self.second_window, width=150, font=('Arial', 18, 'bold'))
        self.combo2.set('')
        self.combo2.place(x=210, y=25)

        self.start_button = ctk.CTkButton(self.second_window, width=60, text="Підключитись", command=lambda: self.connect_journal(self.combo, self.combo2), font=('Arial', 18, 'bold'))
        self.start_button.place(x=120, y=80)

        self.file_combo = ctk.CTkComboBox(self.second_window, width=200, state="readonly")

        self.combo.focus_set()
    
    def update_file_list(self):
        initial_directory = "/путь/к/начальной/папке"
        self.file_path = filedialog.askopenfilename(initialdir=initial_directory, filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")])

        
        if self.file_path:
            self.file_combo.set(self.file_path)

    def connect_journal(self, combo, combo2):
        desired_subject = combo.get()
        desired_class = combo2.get()
        subject_elements = self.driver.find_elements('xpath', "//td[not(a[contains(@class, 'gray-button-2')])]")
        for element in subject_elements:
            subject_name = element.text.strip() 
            if subject_name == desired_subject:
                class_links = element.find_elements('xpath', "./following-sibling::td/a[@class='gray-button-2']")
                for class_link in class_links:
                    class_text = class_link.text
                    if desired_class in class_text:
                        class_link.click()
                        self.driver.sleep(2)
                        self.hide_widgets()
                        break
                break



    def journal(self):
        elements = self.driver.find_elements('xpath', "//table[@class='journal-choose']//td[not(a[contains(@class, 'gray-button-2')])]")
        subjects = [element.text.strip() for element in elements if element.text.strip()]
        return subjects
    
        
    def hide_widgets(self):
        self.combo.place_forget()
        self.combo2.place_forget()
        self.start_button.place_forget()
        self.second_window.geometry('500x400+200+100')
        self.back_button = ctk.CTkButton(self.second_window, width=80,text='Журнали', command=self.back_journals , font=('Arial', 16, 'bold'))
        self.back_button.place(x=17, y=20)
        
        self.file_select_button = ctk.CTkButton(self.second_window, text='Виберіть файл', command=lambda:self.update_file_list(), font=('Arial', 18, 'bold'))
        self.file_select_button.place(x=17, y=50)
        
        self.new_row_label = ctk.CTkLabel(self.second_window, text="Виберіть номер рядка:")
        self.new_row_label.place(x=60, y=100)
        self.new_row_entry = ctk.CTkEntry(self.second_window, width=50)
        self.new_row_entry.place(x=250, y=100)
        self.new_column_label = ctk.CTkLabel(self.second_window, text="Виберіть номер стовпчика:")
        self.new_column_label.place(x=60, y=140)
        self.new_column_entry = ctk.CTkEntry(self.second_window, width=50)
        self.new_column_entry.place(x=250, y=140)

        self.seq_label = ctk.CTkLabel(self.second_window, text="Виберіть номер урока:")
        self.seq_label.place(x=60, y=180)
        self.seq_entry = ctk.CTkEntry(self.second_window, width=50)
        self.seq_entry.place(x=250, y=180)

        self.doc_button=ctk.CTkButton(self.second_window, text='Запустити', command=self.docx)
        self.doc_button.place(x=209, y=250)

        self.file_combo.place (x=175,y=50)
    

    def back_journals(self):
        self.driver.open('https://nz.ua/journal/list')
        self.second_window.withdraw()
        self.open_second_window()
        self.combo.focus_set()

    def docx(self):
        document_path = self.file_path
        doc = Document(document_path)
        table = doc.tables[0]
        chan = self.driver.find_elements(".dz-edit")
        row_index = int(self.new_row_entry.get())
        lesson_number = int(self.seq_entry.get())
        column_index = int(self.new_column_entry.get())
        self.second_window.withdraw()

        save_button_text = "Зберегти"
        
        # empty_lessons_indices = []

        for i in range(len(chan)):
            print(i)
            leson = self.driver.find_elements('.dzc-theme')
            self.driver.sleep(1)
            if not leson[i].text.strip():
                # empty_lessons_indices.append(i)
                chan = self.driver.find_elements(".dz-edit")
                element = chan[i]
                element.click()
                leson = self.driver.find_elements('.dzc-theme')
                print(row_index)
                cell_text = table.cell(row_index - 1 + i, column_index - 1).text
                self.driver.type('#osvitaschedulereal-lesson_topic', cell_text)
                self.driver.type('#osvitaschedulereal-lesson_number_in_plan', str(lesson_number+ i))
                self.driver.sleep(1)
                self.driver.click_link(save_button_text)
                self.driver.sleep(1)
            
            else:
                print (f"Урок {i+1} заповнений!")

        self.second_window.deiconify()

if __name__ == '__main__':
    app = WebAutomationApp()
